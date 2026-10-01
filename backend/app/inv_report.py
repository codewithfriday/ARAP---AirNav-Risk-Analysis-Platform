"""Investigation report generator for ATSB-method (ORLIO) safety investigation studies.

One structured builder (`build_report`) produces the report as a list of sections made of simple blocks; the
same structure is rendered as JSON (editor preview), DOCX (python-docx) and PDF (reportlab). Report structure:

1. Executive summary (synopsis, consequences, core finding)
2. Factual information (timeline, personnel and assets, immediate actions)
3. ORLIO analysis framework (3.1 occurrence events · 3.2 individual actions · 3.3 local conditions ·
   3.4 risk controls · 3.5 organisational influences)
4. Findings and contributing factors (prioritised chain of contributing factors)
5. Safety recommendations and corrective actions (hierarchy of controls)
6. Appendices (A scene photographs/diagrams · B interviews · C ORLIO factor map)
"""
from __future__ import annotations

import base64
import io
import re
import textwrap
from datetime import date
from typing import Any

from . import atsb as ref
from .engines import atsb as engine

REPORT_VERSION = "inv-report-1.0.0"
LAYER_OF = {t: k for k, _, _, ts, _ in ref.REPORT_LAYERS for t in ts}
LAYER_NAME = {k: nm for k, _, nm, _, _ in ref.REPORT_LAYERS}
LAYER_SHORT = {"E": "Occurrence event", "I": "Individual action", "L": "Local condition", "R": "Risk control", "O": "Organisational influence"}
LAYER_COL = {"E": "Occurrence events", "I": "Individual actions", "L": "Local conditions", "R": "Risk controls", "O": "Organisational"}
TYPE_SHORT = {"OE": "Occurrence event", "TFM": "Technical failure", "IA": "Individual action", "PA": "Positive action",
              "LC": "Local condition", "RC": "Risk control", "PC": "Positive condition", "OI": "Organisational influence"}
HIER = {k: n for k, n, _ in ref.HIERARCHY}
HIER_RANK = {k: i for i, (k, _, _) in enumerate(ref.HIERARCHY)}
PROB = {c: t for c, t, _, _ in ref.PROBABILITY}
FINDING_LABEL = {"contributing": "Contributing safety factor", "other": "Other factor that increased risk", "positive": "Positive factor",
                 "not_established": "Not established to the standard of proof", "not_safety_factor": "Not a safety factor",
                 "pending": "Analysis pending", "excluded": "Not analysed further"}
LEVEL_RANK = {"critical": 3, "significant": 2, "minor": 1, "broadly_acceptable": 0, None: 0}
# ORLIO map colours — identical to the editor's AcciMap
FILL = {"O": "#F4B6D2", "R": "#CDBAF0", "L": "#FBE67A", "I": "#FBC477", "T": "#F6D7A7", "E": "#B7DBA3"}
LANE_BG = {"O": "#FDF0F6", "R": "#F5F0FC", "L": "#FFFBE3", "I": "#FFF4E3", "E": "#F1F8ED"}
NAVY = "#1F3A5F"


# ------------------------------------------------------------------ helpers
def _p(text, style=None):
    return {"type": "p", "text": text or "—", **({"style": style} if style else {})}


def _h(text, level=2):
    return {"type": "h", "text": text, "level": level}


def _table(header, rows, widths=None):
    return {"type": "table", "header": header, "rows": [["" if v is None else str(v) for v in r] for r in rows],
            **({"widths": widths} if widths else {})}


def _bullets(items):
    return {"type": "bullets", "items": [i for i in items if i]}


def _kv(rows):
    return {"type": "kv", "rows": [[k, v] for k, v in rows if v]}


def _fmt_time(e: dict) -> str:
    s = e.get("start") or ""
    s = s.replace("T", " ")
    return s + (f" – {e['end'].replace('T', ' ')}" if e.get("end") else "")


def _ids(fid: str) -> tuple:
    m = re.match(r"([A-Za-z]*)(\d+)", fid or "")
    return (m.group(1), int(m.group(2))) if m else (fid, 0)


def _layer(f: dict) -> str | None:
    return LAYER_OF.get(f.get("type"))


def _chain(fid: str, by_id: dict, limit: int = 12) -> list[str]:
    out, cur, seen = [fid], fid, {fid}
    while len(out) < limit:
        tg = ((by_id.get(cur) or {}).get("influence") or {}).get("target")
        if not tg or tg in seen:
            break
        out.append(tg)
        if tg == "occurrence":
            break
        seen.add(tg); cur = tg
    return out


def _short(f: dict | None, fid: str) -> str:
    if fid == "occurrence":
        return "the occurrence"
    return f"{fid} ({TYPE_SHORT.get((f or {}).get('type'), (f or {}).get('type', ''))})"


def _evidence(test: dict, n: int = 4) -> list[str]:
    out = []
    for it in (test or {}).get("items") or []:
        r = it.get("rating")
        if r not in ("supports", "opposes"):
            continue
        tag = "supports" if r == "supports" else "opposes"
        out.append(f"[{tag}] {it.get('text', '')}" + (f" — {it['source']}" if it.get("source") else ""))
        if len(out) >= n:
            break
    return out


def _test_line(test: dict) -> str:
    if not test:
        return ""
    c = test.get("conclusion") or "pending"
    pr = PROB.get(test.get("probability") or "", "")
    return f"{c}" + (f" ({pr})" if pr else "")


def _owner(a: dict) -> str:
    return a.get("organisation") or "[Owner]"


# ------------------------------------------------------------------ builder
def build_report(model: dict, analysis: dict | None = None, project: dict | None = None, airnav: dict | None = None) -> dict[str, Any]:
    analysis = analysis or engine.analyse(model, airnav)
    occ = model.get("occurrence") or {}
    rep = model.get("report") or {}
    factors = [f for f in model.get("factors") or []]
    by_id = {f["id"]: f for f in factors}
    res = analysis.get("factors") or {}
    ft = {fid: (res.get(fid) or {}).get("finding_type") for fid in by_id}
    active = [f for f in factors if ft.get(f["id"]) in ("contributing", "other", "positive", "pending")]
    contributing = [f for f in factors if ft.get(f["id"]) == "contributing"]
    issues = {i["id"]: i for i in analysis.get("issues") or []}

    def title(fid):
        return "the occurrence" if fid == "occurrence" else (by_id.get(fid) or {}).get("title", fid)

    sections = []

    # ---------------------------------------------------------- 1 executive summary
    s1 = []
    s1.append(_h("1.1 Synopsis"))
    s1.append(_kv([("What", occ.get("title") or occ.get("occurrence_type")), ("Occurrence type", occ.get("occurrence_type")),
                   ("When", " ".join(x for x in (occ.get("date"), occ.get("time")) if x)), ("Where", occ.get("location")),
                   ("Who", occ.get("operator")), ("Reference", occ.get("ref"))]))
    if occ.get("summary"):
        s1.append(_p(occ["summary"]))
    s1.append(_h("1.2 Consequences"))
    s1.append(_p(rep.get("consequences") or "Injuries, damage and environmental impact not yet recorded."))
    s1.append(_h("1.3 Core finding"))
    if rep.get("core_finding"):
        s1.append(_p(rep["core_finding"]))
    by_layer = {k: [f for f in contributing if _layer(f) == k] for k in LAYER_NAME}
    lvl_counts: dict[str, int] = {}
    for i in issues.values():
        lvl_counts[i["level"]] = lvl_counts.get(i["level"], 0) + 1
    auto = []
    oe = by_layer["E"]
    if oe:
        auto.append("Occurrence: " + "; ".join(f["title"] for f in oe) + ".")
    for k, lead in (("I", "Individual actions"), ("L", "Local conditions"), ("R", "Risk controls that were absent, failed or were bypassed"),
                    ("O", "Organisational influences")):
        if by_layer[k]:
            auto.append(f"{lead}: " + "; ".join(f["title"].rstrip(".") for f in by_layer[k]) + ".")
    auto_p = (f"The ORLIO analysis identified {len(contributing)} contributing safety factors and "
              f"{sum(1 for v in ft.values() if v == 'other')} other factors that increased risk. ")
    if issues:
        auto_p += (f"{len(issues)} safety issues were identified ("
                   + ", ".join(f"{n} {lvl.replace('_', ' ')}" for lvl, n in sorted(lvl_counts.items(), key=lambda x: -LEVEL_RANK.get(x[0], 0)))
                   + "). ")
    auto_p += "The primary safety breakdowns, traced from the occurrence back to the organisational level, were:"
    s1.append(_p(auto_p))
    s1.append(_bullets(auto))
    s1.append(_p("Findings are stated to identify safety factors, not to apportion blame or liability.", "note"))
    sections.append({"id": "s1", "title": "1. Executive summary", "blocks": s1})

    # ---------------------------------------------------------- 2 factual information
    s2 = [_h("2.1 Occurrence timeline")]
    evs = [e for e in model.get("events") or [] if e.get("display", True)]
    if evs:
        s2.append(_table(["Time", "Event", "Source"], [[_fmt_time(e), e.get("title", ""), e.get("source", "")] for e in evs], [18, 62, 20]))
    else:
        s2.append(_p("No sequence of events recorded."))
    s2.append(_h("2.2 Personnel and assets involved"))
    if rep.get("personnel"):
        s2.append(_table(["Personnel", "Experience, qualifications and duty"], [[p.get("role"), p.get("details")] for p in rep["personnel"]], [28, 72]))
    if rep.get("assets"):
        s2.append(_table(["Asset / equipment", "Type and configuration"], [[a.get("item"), a.get("details")] for a in rep["assets"]], [28, 72]))
    if rep.get("environment"):
        s2.append(_kv([("Environmental configuration", rep["environment"])]))
    if not (rep.get("personnel") or rep.get("assets") or rep.get("environment")):
        s2.append(_p("Not yet recorded."))
    s2.append(_h("2.3 Immediate actions taken"))
    s2.append(_bullets(rep.get("immediate_actions") or []) if rep.get("immediate_actions") else _p("Not yet recorded."))
    sections.append({"id": "s2", "title": "2. Factual information", "blocks": s2})

    # ---------------------------------------------------------- 3 ORLIO analysis
    s3 = [_p("Evidence is mapped against the ORLIO layers to trace the causal flow from the occurrence back to organisational "
             "influences. Each factor was tested for existence, and for influence on the occurrence or on another factor "
             f"(or for importance), to the standard of proof '{PROB['L'].lower()}' (≥{ref.STANDARD_OF_PROOF} %), following the ATSB "
             "safety investigation analysis method.")]
    notes = rep.get("layer_notes") or {}
    for key, num, name, types, definition in ref.REPORT_LAYERS:
        s3.append(_h(f"{num} {name}"))
        s3.append(_p(f"Definition: {definition}", "definition"))
        if notes.get(key):
            s3.append(_p(notes[key]))
        lay = [f for f in active if f.get("type") in types]
        lay.sort(key=lambda f: (0 if ft[f["id"]] == "contributing" else 1, _ids(f["id"])))
        if not lay:
            s3.append(_p("No factors at this level were established."))
            continue
        for f in lay:
            r = res.get(f["id"]) or {}
            s3.append(_h(f"{f['id']} · {f['title']}", 3))
            tm = ref.TYPES.get(f.get("type"), {})
            inf = f.get("influence") or {}
            rows = [("Classification", f"{tm.get('name', f.get('type'))}" + (f" · {', '.join(f.get('codes') or [])}" if f.get("codes") else "")),
                    ("Finding", FINDING_LABEL.get(ft[f["id"]], ft[f["id"]])),
                    ("Description", f.get("description")),
                    ("Role / error type", " · ".join(x for x in (f.get("role"), f.get("error_type")) if x) if f.get("type") in ("IA", "PA") else ""),
                    ("Why it made sense at the time", f.get("rationale") if f.get("type") in ("IA", "PA") else ""),
                    ("Control function", f.get("control_function") if f.get("type") in ("RC", "PC") else ""),
                    ("Existence", " — ".join(x for x in (_test_line(f.get("existence") or {}), (f.get("existence") or {}).get("summary")) if x)),
                    ("Influence", " — ".join(x for x in (_test_line(inf), f"on {_short(by_id.get(inf.get('target')), inf['target'])}: {title(inf['target'])}"
                                                          if inf.get("target") else "", inf.get("summary")) if x)),
                    ("Importance", (f.get("importance") or {}).get("justification") if r.get("importance") is not None else ""),
                    ("Explained by", ", ".join(f"{x} {title(x)}" for x in r.get("explained_by") or [])),
                    ("Further explanation", f.get("sufficiency_note"))]
            if f["id"] in issues:
                i = issues[f["id"]]
                rows.append(("Safety issue", f"{i['risk'].get('region_name', '')} — risk {i['risk'].get('index', '')} "
                                             f"({'ATSB 6×6' if i['risk'].get('scheme') == 'atsb' else 'AirNav 5×5'}); owner {i.get('owner') or '—'}"))
            s3.append(_kv(rows))
            ev = _evidence(f.get("existence") or {}) + _evidence(inf, 2)
            if ev:
                s3.append({"type": "bullets", "items": ev, "style": "evidence"})
    sections.append({"id": "s3", "title": "3. ORLIO analysis framework", "blocks": s3})

    # ---------------------------------------------------------- 4 findings
    order = {k: i for i, k in enumerate(["O", "R", "L", "I", "E"])}

    def prio(f):
        i = issues.get(f["id"])
        return (-LEVEL_RANK.get(i["level"] if i else None, 0), -(i["risk"].get("rank", 0) if i else 0), order.get(_layer(f), 9), _ids(f["id"]))

    s4 = [_p("Contributing safety factors are listed in priority order — safety issues by risk level first, then from the organisational "
             "level down to the occurrence. Each entry shows the chain through which it contributed (→ influenced).")]
    s4.append(_h("4.1 Contributing factors (chain of events and conditions)"))
    if oe:
        s4.append(_p("Occurrence: " + "; ".join(f"{f['id']} {f['title']}" for f in oe) + "."))
    items = []
    for n, f in enumerate(sorted([f for f in contributing if f.get("type") != "OE"], key=prio), 1):
        ch = _chain(f["id"], by_id)
        path = " → ".join(_short(by_id.get(x), x) for x in ch)
        si = issues.get(f["id"])
        tag = f" [Safety issue · {si['risk'].get('region_name', '')} {si['risk'].get('index', '')}]" if si else ""
        items.append(f"Contributing factor {n} — {TYPE_SHORT.get(f.get('type'), '')}: {f['title']}.{tag} Chain: {path}.")
    s4.append({"type": "numbered", "items": items} if items else _p("No contributing factors established."))
    oth = [f for f in factors if ft.get(f["id"]) == "other"]
    s4.append(_h("4.2 Other factors that increased risk"))
    s4.append({"type": "numbered", "items": [f"{TYPE_SHORT.get(f.get('type'), '')}: {f['title']}." + (" [Safety issue]" if f["id"] in issues else "")
                                            for f in sorted(oth, key=prio)]} if oth else _p("None."))
    pos = [f for f in factors if ft.get(f["id"]) == "positive"]
    if pos:
        s4.append(_h("4.3 Positive factors"))
        s4.append(_bullets([f["title"] for f in pos]))
    kfs = [k for k in model.get("key_findings") or [] if (analysis.get("key_findings") or {}).get(k.get("id"), {}).get("supported")]
    s4.append(_h("4.4 Other key findings"))
    s4.append(_bullets([k.get("statement") for k in kfs]) if kfs else _p("None."))
    if issues:
        s4.append(_h("4.5 Safety issues"))
        s4.append(_table(["Ref", "Safety issue", "Owner", "Risk", "Status"],
                         [[i["id"], i["title"], i.get("owner") or "—", f"{i['risk'].get('index', '')} {i['risk'].get('region_name', '')}",
                           ref.ISSUE_STATUS.get(i.get("status") or "pending", i.get("status"))] for i in sorted(issues.values(), key=lambda i: -LEVEL_RANK.get(i["level"], 0))],
                         [7, 45, 18, 18, 12]))
    exc = [f for f in factors if ft.get(f["id"]) in ("excluded", "not_safety_factor", "not_established")]
    if exc:
        s4.append(_h("4.6 Factors considered and not included"))
        s4.append(_table(["Ref", "Factor considered", "Reason"],
                         [[f["id"], f["title"], " — ".join(x for x in (f.get("exclusion_reason") or FINDING_LABEL.get(ft[f["id"]]),
                                                                           f.get("further_justification")) if x)]
                          for f in exc], [7, 53, 40]))
    sections.append({"id": "s4", "title": "4. Findings and contributing factors", "blocks": s4})

    # ---------------------------------------------------------- 5 corrective actions
    s5 = [_p("Every identified contributing factor is resolved with corrective action, ranked by the hierarchy of controls: "
             + " > ".join(n for _, n, _ in ref.HIERARCHY) + ". Formal safety recommendations state the safety issue; the "
             "addressee determines the means. [Date] marks a target date still to be agreed.")]
    rows, flagged = [], 0
    for f in sorted([f for f in active if f.get("actions")], key=lambda f: (order.get(_layer(f), 9), _ids(f["id"]))) + \
            [f for f in contributing if not f.get("actions") and f.get("type") != "OE"]:
        layer = LAYER_COL.get(_layer(f), "")
        if not f.get("actions"):
            flagged += 1
            rows.append([layer, f"{f['id']} {f['title']}", "[No corrective action recorded — to be determined]", "[Owner]", "[Date]"])
            continue
        for a in sorted(f["actions"], key=lambda a: HIER_RANK.get(a.get("hierarchy"), 9)):
            h = HIER.get(a.get("hierarchy"))
            kind = {"recommendation": "Safety recommendation", "advisory": "Safety advisory notice"}.get(a.get("kind"), "")
            text = (f"[{h}] " if h else "") + (f"{kind}: " if kind else "") + (a.get("description") or "")
            extra = "; ".join(x for x in (a.get("ref"), ref.ACTION_STATUS.get(a.get("status"), "")) if x)
            rows.append([layer, f"{f['id']} {f['title']}", text + (f" ({extra})" if extra else ""), _owner(a),
                         a.get("target_date") or ("Completed" if a.get("status") == "closed" else "[Date]")])
    s5.append(_table(["ORLIO layer", "Identified deficiency", "Corrective action required", "Action owner", "Target date"], rows,
                     [12, 26, 36, 14, 12]) if rows else _p("No corrective actions recorded."))
    hc: dict[str, int] = {}
    for f in active:
        for a in f.get("actions") or []:
            hc[a.get("hierarchy") or "unclassified"] = hc.get(a.get("hierarchy") or "unclassified", 0) + 1
    if hc:
        s5.append(_p("Actions by hierarchy level: " + ", ".join(f"{HIER.get(k, 'Not classified')} {v}" for k, v in
                                                               sorted(hc.items(), key=lambda x: HIER_RANK.get(x[0], 9))) + "."
                     + (f" {flagged} contributing factor(s) still need a corrective action." if flagged else " Every contributing factor has at least one corrective action.")))
    sections.append({"id": "s5", "title": "5. Safety recommendations and corrective actions", "blocks": s5})

    # ---------------------------------------------------------- 6 appendices
    s6 = [_h("Appendix A — Photographs and diagrams of the scene and asset damage")]
    att = [a for a in rep.get("attachments") or [] if (a.get("data") or "").startswith("data:image/")]
    if att:
        for n, a in enumerate(att, 1):
            s6.append({"type": "image", "data": a["data"], "caption": f"Figure A{n}. {a.get('caption') or ''}".strip()})
    else:
        s6.append(_p("No photographs or diagrams attached."))
    s6.append(_h("Appendix B — Interview summaries and witness accounts"))
    iv = rep.get("interviews") or []
    if iv:
        for x in iv:
            s6.append({"type": "kv", "rows": [[x.get("person") or "Interviewee", x.get("summary") or ""]]})
    else:
        s6.append(_p("No interview summaries recorded."))
    s6.append(_h("Appendix C — ORLIO factor map"))
    s6.append({"type": "image", "data": "data:image/png;base64," + base64.b64encode(factor_map_png(model, analysis)).decode(),
               "caption": "Figure C1. ORLIO factor map — solid borders are contributing safety factors, dashed borders other factors that "
                          "increased risk; solid arrows show an established influence, dotted arrows a relationship considered but not "
                          "established; SI badges mark safety issues with their risk index."})
    sections.append({"id": "s6", "title": "6. Appendices and supporting evidence", "blocks": s6})

    return {"title": "Safety Investigation Report", "subtitle": occ.get("title") or "",
            "meta": {"report_no": rep.get("report_no") or (f"{project['code']}-INV" if project else ""),
                     "prepared_by": rep.get("prepared_by") or "", "status": rep.get("status") or "Draft",
                     "project": f"{project['code']} — {project['title']}" if project else "",
                     "generated": date.today().isoformat(), "engine": analysis.get("engine_version"), "report_version": REPORT_VERSION,
                     "source": rep.get("source") or ""},
            "sections": sections}


# ------------------------------------------------------------------ factor map (Appendix C)
def factor_map_png(model: dict, analysis: dict, dpi: int = 200) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle

    res = analysis.get("factors") or {}
    issues = {i["id"]: i for i in analysis.get("issues") or []}
    fs = [f for f in model.get("factors") or [] if (res.get(f["id"]) or {}).get("finding_type") in ("contributing", "other", "positive")]
    lanes = ["O", "R", "L", "I", "E"]
    lane_of = {f["id"]: _layer(f) or "E" for f in fs}
    groups = {k: sorted([f for f in fs if lane_of[f["id"]] == k], key=lambda f: _ids(f["id"])) for k in lanes}
    W = max(12.0, 2.45 * max([len(g) for g in groups.values()] + [1]))
    links = []
    for f in fs:
        tg = (f.get("influence") or {}).get("target")
        if tg and (tg in lane_of or tg == "occurrence"):
            links.append((f["id"], tg, bool((res.get(f["id"]) or {}).get("influence"))))
    pos: dict[str, float] = {}

    def spread(k):
        n = len(groups[k])
        for i, f in enumerate(groups[k]):
            pos[f["id"]] = (i + 0.5) * W / max(n, 1)
    for k in lanes:
        spread(k)
    oe = [f["id"] for f in groups["E"] if f.get("type") == "OE"]
    for _ in range(8):  # barycentre sweeps to reduce crossings
        pos["occurrence"] = pos[oe[0]] if oe else W / 2
        for k in lanes:
            def bc(f):
                nb = [pos[b] for a, b, _ in links if a == f["id"] and b in pos] + [pos[a] for a, b, _ in links if b == f["id"]]
                return sum(nb) / len(nb) if nb else pos[f["id"]]
            groups[k].sort(key=bc)
            spread(k)
    pos["occurrence"] = pos[oe[0]] if oe else W / 2
    LH, BH, S = 1.8, 1.5, 0.52
    y_of = {k: (len(lanes) - 1 - i) * LH for i, k in enumerate(lanes)}
    top = y_of["O"] + LH / 2
    ybot = -LH / 2 - 1.05
    fig = plt.figure(figsize=((W + 1.6) * S, (top - ybot) * S))
    ax = fig.add_axes((0, 0, 1, 1))
    ax.set_xlim(-1.5, W + 0.1); ax.set_ylim(ybot, top); ax.axis("off")
    for k in lanes:
        y = y_of[k]
        ax.add_patch(Rectangle((-1.5, y - LH / 2 + 0.04), W + 1.6, LH - 0.08, fc=LANE_BG[k], ec="#e5e7eb", lw=0.5))
        ax.text(-1.42, y, LAYER_NAME[k].replace(" (", "\n(").replace("Organisational ", "Organisational\n"), ha="left", va="center",
                fontsize=6, weight="bold", color="#374151")
    for k in lanes:
        n = len(groups[k])
        bw = min(2.25, W / max(n, 1) - 0.18)
        for f in groups[k]:
            fid = f["id"]; x, y = pos[fid], y_of[k]
            fty = res[fid]["finding_type"]
            contrib = fty == "contributing"
            fill = FILL["T"] if f.get("type") == "TFM" else FILL[k]
            ax.add_patch(FancyBboxPatch((x - bw / 2, y - BH / 2), bw, BH, boxstyle="round,pad=0,rounding_size=0.07", fc=fill,
                                        ec=NAVY if contrib else "#6B7280", lw=1.2 if contrib else 0.8, ls="-" if contrib else (0, (3, 2)), zorder=2))
            chars = max(14, int(bw * 11.5))
            lines = textwrap.wrap(f.get("title", ""), chars)
            if len(lines) > 6:
                lines = lines[:6]; lines[-1] = lines[-1][: chars - 1].rstrip() + "…"
            head = f"{fid}" + (f" · {', '.join((f.get('codes') or [])[:2])}" if f.get("codes") else "")
            ax.text(x - bw / 2 + 0.07, y + BH / 2 - 0.07, head, ha="left", va="top", fontsize=5.6, weight="bold", color=NAVY, zorder=3)
            ax.text(x, y - 0.12, "\n".join(lines), ha="center", va="center", fontsize=5.4, color="#111827", linespacing=1.08, zorder=3)
            if fid in issues:
                r = issues[fid]["risk"]
                ax.add_patch(FancyBboxPatch((x + bw / 2 - 0.62, y + BH / 2 - 0.13), 0.58, 0.26, boxstyle="round,pad=0,rounding_size=0.05",
                                            fc=r.get("color", "#E3B23C"), ec="white", lw=0.5, zorder=4))
                ax.text(x + bw / 2 - 0.33, y + BH / 2, f"SI {r.get('index', '')}", ha="center", va="center", fontsize=5, color="white",
                        weight="bold", zorder=5)
    yo = y_of["E"] - LH / 2 - 0.3
    xo = pos["occurrence"]
    ax.add_patch(FancyBboxPatch((xo - 0.9, yo - 0.16), 1.8, 0.32, boxstyle="round,pad=0,rounding_size=0.1", fc=NAVY, ec=NAVY, zorder=2))
    ax.text(xo, yo, "OCCURRENCE", ha="center", va="center", fontsize=6, color="white", weight="bold", zorder=3)
    bwd = {f["id"]: min(2.25, W / max(len(groups[lane_of[f["id"]]]), 1) - 0.18) for f in fs}
    for a, b, est in links:
        xa, ya = pos[a], y_of[lane_of[a]]
        contrib = res[a]["finding_type"] == "contributing"
        col = NAVY if (contrib and est) else "#6B7280" if est else "#9CA3AF"
        st = dict(arrowstyle="-|>", mutation_scale=7, lw=0.9 if contrib else 0.7, color=col, ls="-" if est else (0, (2, 2)), zorder=1,
                  shrinkA=0, shrinkB=0)
        if b == "occurrence":
            ax.add_patch(FancyArrowPatch((xa, ya - BH / 2), (xo, yo + 0.16), **st)); continue
        xb, yb = pos[b], y_of[lane_of[b]]
        if abs(yb - ya) < 1e-6:  # same lane: side to side
            sgn = 1 if xb > xa else -1
            between = any(min(xa, xb) < pos[g["id"]] < max(xa, xb) for g in groups[lane_of[a]])
            if between:  # arc over the boxes in between
                ax.add_patch(FancyArrowPatch((xa, ya + BH / 2), (xb, yb + BH / 2), connectionstyle=f"arc3,rad={0.25 * sgn}", **{**st, "zorder": 6}))
            else:
                ax.add_patch(FancyArrowPatch((xa + sgn * bwd[a] / 2, ya), (xb - sgn * bwd[b] / 2, yb), **st))
        elif yb < ya:
            ax.add_patch(FancyArrowPatch((xa, ya - BH / 2), (xb, yb + BH / 2), **st))
        else:
            ax.add_patch(FancyArrowPatch((xa, ya + BH / 2), (xb, yb - BH / 2), connectionstyle="arc3,rad=0.15", **st))
    handles = [Line2D([], [], color=NAVY, lw=1.2, label="Contributing safety factor / influence established"),
               Line2D([], [], color="#6B7280", lw=0.8, ls=(0, (3, 2)), label="Other factor that increased risk"),
               Line2D([], [], color="#9CA3AF", lw=0.8, ls=(0, (2, 2)), label="Relationship considered; influence not established")]
    ax.legend(handles=handles, loc="lower right", bbox_to_anchor=(1.0, 0.0), ncol=1, fontsize=5.4, frameon=False, handlelength=2.4)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


# ------------------------------------------------------------------ renderers
def _img_bytes(data_url: str) -> bytes:
    return base64.b64decode(data_url.split(",", 1)[1])


def to_docx(report: dict) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    navy = RGBColor(0x1F, 0x3A, 0x5F)
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    h = doc.add_heading(report["title"], 0); h.runs[0].font.color.rgb = navy
    if report.get("subtitle"):
        p = doc.add_paragraph(); r = p.add_run(report["subtitle"]); r.bold = True; r.font.size = Pt(12)
    m = report["meta"]
    p = doc.add_paragraph()
    for label, key in (("Report", "report_no"), ("Project", "project"), ("Status", "status"), ("Prepared by", "prepared_by"),
                       ("Source", "source"), ("Generated", "generated")):
        if m.get(key):
            p.add_run(f"{label}: ").bold = True; p.add_run(f"{m[key]}\n")
    p.add_run(f"NAVRAP {m.get('report_version')} · analysis engine {m.get('engine')}").italic = True
    usable = 17.0
    for sec in report["sections"]:
        doc.add_page_break() if sec["id"] == "s6" else None
        hh = doc.add_heading(sec["title"], 1); hh.runs[0].font.color.rgb = navy
        for b in sec["blocks"]:
            t = b["type"]
            if t == "h":
                x = doc.add_heading(b["text"], b.get("level", 2)); x.runs[0].font.color.rgb = navy
            elif t == "p":
                x = doc.add_paragraph()
                r = x.add_run(b["text"])
                if b.get("style") in ("note", "definition"):
                    r.italic = True
                    if b.get("style") == "note":
                        r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            elif t in ("bullets", "numbered"):
                for it in b["items"]:
                    x = doc.add_paragraph(it, style="List Number" if t == "numbered" else "List Bullet")
                    if b.get("style") == "evidence":
                        for r in x.runs:
                            r.font.size = Pt(8.5)
            elif t == "kv":
                tb = doc.add_table(rows=0, cols=2); tb.style = "Table Grid"
                for k, v in b["rows"]:
                    c = tb.add_row().cells
                    c[0].text = ""; c[0].paragraphs[0].add_run(str(k)).bold = True; c[1].text = str(v)
                    c[0].width = Cm(4.2); c[1].width = Cm(usable - 4.2)
                doc.add_paragraph()
            elif t == "table":
                tb = doc.add_table(rows=1, cols=len(b["header"])); tb.style = "Light Grid Accent 1"
                for i, hd in enumerate(b["header"]):
                    tb.rows[0].cells[i].text = hd
                for row in b["rows"]:
                    c = tb.add_row().cells
                    for i, v in enumerate(row):
                        c[i].text = v
                if b.get("widths"):
                    tot = sum(b["widths"])
                    for row in tb.rows:
                        for i, w in enumerate(b["widths"]):
                            row.cells[i].width = Cm(usable * w / tot)
                for row in tb.rows:
                    for c in row.cells:
                        for pp in c.paragraphs:
                            for r in pp.runs:
                                r.font.size = Pt(8.5)
                doc.add_paragraph()
            elif t == "image":
                doc.add_picture(io.BytesIO(_img_bytes(b["data"])), width=Cm(usable))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp = doc.add_paragraph(); r = cp.add_run(b.get("caption", "")); r.italic = True; r.font.size = Pt(8.5)
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


def to_pdf(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)
    from xml.sax.saxutils import escape

    navy = colors.HexColor(NAVY)
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9, leading=12, alignment=TA_LEFT, spaceAfter=4)
    small = ParagraphStyle("s", parent=body, fontSize=7.8, leading=9.6, spaceAfter=0)
    small_b = ParagraphStyle("sb", parent=small, fontName="Helvetica-Bold")
    note = ParagraphStyle("n", parent=body, fontName="Helvetica-Oblique", textColor=colors.HexColor("#6B7280"))
    defin = ParagraphStyle("d", parent=body, fontName="Helvetica-Oblique")
    H = {0: ParagraphStyle("h0", parent=body, fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=navy, spaceAfter=6),
         1: ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=navy, spaceBefore=8, spaceAfter=6),
         2: ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=navy, spaceBefore=6, spaceAfter=4),
         3: ParagraphStyle("h3", parent=body, fontName="Helvetica-Bold", fontSize=9.4, leading=12, textColor=colors.HexColor("#111827"),
                           spaceBefore=5, spaceAfter=3)}
    cap = ParagraphStyle("c", parent=small, fontName="Helvetica-Oblique", spaceBefore=3, spaceAfter=8)
    W = A4[0] - 36 * mm

    def P(t, s=body):
        return Paragraph(escape(str(t)).replace("\n", "<br/>"), s)

    m = report["meta"]
    story = [P(report["title"], H[0])]
    if report.get("subtitle"):
        story.append(P(report["subtitle"], H[2]))
    meta_rows = [[P(k, small_b), P(m[key], small)] for k, key in (("Report", "report_no"), ("Project", "project"), ("Status", "status"),
                                                                  ("Prepared by", "prepared_by"), ("Source", "source"), ("Generated", "generated")) if m.get(key)]
    if meta_rows:
        t = Table(meta_rows, colWidths=[30 * mm, W - 30 * mm])
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#e5e7eb"))]))
        story += [t, Spacer(1, 3 * mm)]
    story.append(P(f"NAVRAP {m.get('report_version')} · analysis engine {m.get('engine')}", note))
    for sec in report["sections"]:
        if sec["id"] == "s6":
            story.append(PageBreak())
        story.append(P(sec["title"], H[1]))
        pending_head = None
        for b in sec["blocks"]:
            t = b["type"]
            fl = []
            if t == "h":
                pending_head = P(b["text"], H.get(b.get("level", 2) + (0 if b.get("level", 2) <= 3 else -1), H[3]))
                continue
            if t == "p":
                fl = [P(b["text"], note if b.get("style") == "note" else defin if b.get("style") == "definition" else body)]
            elif t in ("bullets", "numbered"):
                st = small if b.get("style") == "evidence" else body
                fl = [ListFlowable([ListItem(P(i, st), leftIndent=12) for i in b["items"]],
                                   bulletType="1" if t == "numbered" else "bullet", start="1" if t == "numbered" else None,
                                   bulletFontSize=8, leftIndent=12)]
            elif t == "kv":
                tb = Table([[P(k, small_b), P(v, small)] for k, v in b["rows"]], colWidths=[38 * mm, W - 38 * mm])
                tb.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#d1d5db")),
                                        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3F4F6"))]))
                fl = [tb, Spacer(1, 2 * mm)]
            elif t == "table":
                ws = b.get("widths") or [1] * len(b["header"])
                tot = sum(ws)
                data = [[P(h, ParagraphStyle("th", parent=small_b, textColor=colors.white)) for h in b["header"]]] + \
                       [[P(v, small) for v in r] for r in b["rows"]]
                tb = Table(data, colWidths=[W * w / tot for w in ws], repeatRows=1)
                tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), navy), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#d1d5db")),
                                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FC")])]))
                fl = [tb, Spacer(1, 3 * mm)]
            elif t == "image":
                raw = _img_bytes(b["data"])
                iw, ih = ImageReader(io.BytesIO(raw)).getSize()
                w = W; h = w * ih / iw
                maxh = 200 * mm
                if h > maxh:
                    h = maxh; w = h * iw / ih
                fl = [Image(io.BytesIO(raw), width=w, height=h), P(b.get("caption", ""), cap)]
            if pending_head is not None:
                story.append(KeepTogether([pending_head] + fl[:1])); story += fl[1:]; pending_head = None
            else:
                story += fl

    def foot(c, d):
        c.saveState(); c.setFont("Helvetica", 7); c.setFillColor(colors.HexColor("#6B7280"))
        c.drawString(18 * mm, 10 * mm, f"{m.get('report_no') or 'Safety Investigation Report'} · {m.get('status') or ''}")
        c.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {d.page}")
        c.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                      title=report["title"], author="NAVRAP").build(story, onFirstPage=foot, onLaterPages=foot)
    return buf.getvalue()
