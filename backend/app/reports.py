"""Safety Assessment Report (DOCX) following Manual §32.3 (REP-01, REP-02)."""
import io
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .templates import METHODS

NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def _table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = str(h)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None else str(v)
    return t


def assessment_docx(a: dict, studies: list[dict], engines: dict) -> bytes:
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    h = doc.add_heading("Safety Assessment Report", 0)
    h.runs[0].font.color.rgb = NAVY
    p = doc.add_paragraph()
    p.add_run(f"{a['project']['code']} — {a['project']['title']}\n").bold = True
    p.add_run(f"Assessment: {a['title']} (version {a['version']}, status {a['status']})\n")
    p.add_run(f"Generated {date.today().isoformat()} by NAVRAP. Risk scheme version {a['risk_scheme_version']}. "
              f"Engine versions: {', '.join(f'{k} {v}' for k, v in engines.items())}.")

    doc.add_heading("1. Summary and conclusion", 1)
    worst = a.get("worst_residual_region") or "not yet assessed"
    doc.add_paragraph(f"{len(a['hazards'])} hazards recorded. Worst current/residual risk region: {worst}. "
                      f"Required acceptance authority: {a.get('required_authority') or '—'}.")
    doc.add_heading("2. Description of the change, scope and assumptions", 1)
    for label, key in (("Scope", "scope"), ("Operational environment", "environment"), ("Assumptions", "assumptions")):
        doc.add_paragraph().add_run(label).bold = True
        doc.add_paragraph(a.get(key) or "—")
    doc.add_heading("3. Methods used", 1)
    _table(doc, ["Method", "Study", "Status", "Template"],
           [[METHODS[s["method"]]["name"], s["title"], s["status"], s["template_version"]] for s in studies])
    doc.add_heading("4. Hazards and their assessment", 1)
    _table(doc, ["Ref", "Hazard", "Initial", "Residual", "Region", "Controls"],
           [[x["ref"], x["title"], (x["initial_risk"] or {}).get("index", "—"), (x["residual_risk"] or {}).get("index", "—"),
             ((x["residual_risk"] or x["initial_risk"]) or {}).get("region_name", "—"),
             "; ".join(c["text"] for c in x["controls"])] for x in a["hazards"]])
    doc.add_heading("5. Study results", 1)
    for s in studies:
        doc.add_heading(f"{METHODS[s['method']]['name']}: {s['title']}", 2)
        res = s.get("results") or {}
        m = s.get("model") or {}
        if s["method"] == "fta" and res:
            doc.add_paragraph(f"Top-event probability {res.get('top_probability', 0):.3e}; "
                              f"single points of failure: {', '.join(res.get('single_points_of_failure', [])) or 'none'}.")
            _table(doc, ["Cut set", "Probability", "Contribution"],
                   [[", ".join(c["events"]), f"{c['probability']:.2e}", f"{c['contribution']:.1%}"] for c in res.get("cut_sets", [])[:15]])
        elif s["method"] == "lopa" and res:
            doc.add_paragraph(f"Mitigated frequency {res.get('mitigated_frequency', 0):.2e} {res.get('unit', '')}; "
                              f"target {res.get('target_frequency', 0):.1e}; target met: {res.get('target_met')}.")
        elif s["method"] == "fatigue" and res:
            sp = res.get("samn_perelli") or {}
            if sp.get("n"):
                c = sp["counts"]
                doc.add_paragraph(f"Samn-Perelli pre-shift ratings: {sp['n']} (green {c.get('green', 0)}, caution {c.get('caution', 0)}, "
                                  f"amber {c.get('amber', 0)}, red {c.get('red', 0)}); mean score {sp['mean_score']:.1f}; "
                                  f"{len([x for x in sp.get('checks', []) if x['severity'] != 'info'])} response(s) not meeting the action thresholds.")
            for d in res.get("duties", []):
                if "min_alertness" in d:
                    doc.add_paragraph(f"Duty {d['duty']}: min alertness {d['min_alertness']:.2f} at {d['min_clock']}, "
                                      f"max KSS {d['max_kss']:.2f}, {d['share_at_or_above_threshold']:.0%} of duty at KSS ≥ {d['kss_threshold']}.")
        elif s["method"] == "bbn" and res:
            for k, v in (res.get("marginals") or {}).items():
                doc.add_paragraph(f"P({k}) = " + ", ".join(f"{s_}: {p:.4g}" for s_, p in v.items()))
        elif s["method"] == "bowtie" and m:
            doc.add_paragraph(f"Hazard: {m.get('hazard', '')}. Top event: {m.get('top_event', '')}. "
                              f"{len(m.get('threats', []))} threats, {len(m.get('consequences', []))} consequences, "
                              f"{len(m.get('barriers', {}))} barriers.")
        elif s["method"] == "eta" and res.get("sequences"):
            _table(doc, ["Sequence", "Outcome", "Severity", "Frequency"],
                   [[x["sequence"], x["label"], x.get("severity") or "", f"{x['frequency']:.2e}"] for x in res["sequences"]])
        elif s["method"] == "crm" and res.get("vertical"):
            v = res["vertical"]
            doc.add_paragraph(f"Vertical collision risk {v['total']:.2e} per flight hour against TLS {v['tls']:.1e}: "
                              f"{'meets' if v['meets_tls'] else 'does not meet'} the TLS.")
        elif s["method"] == "hra" and res.get("tasks"):
            _table(doc, ["Task", "Library / GTT", "Nominal HEP", "Assessed HEP"],
                   [[t["id"], f"{t['library'].upper()} {t['gtt']}", f"{t['nominal_hep']:.2g}", f"{t['hep']:.3g}"] for t in res["tasks"]])
        elif s["method"] == "rbd" and res.get("rbd"):
            doc.add_paragraph(f"RBD availability {res['rbd']['availability']:.6f} ({res['rbd']['downtime_hours_per_year']:.2f} h/year downtime). "
                              + (f"Markov unavailability {res['markov']['unavailability']:.2e}." if res.get("markov") else ""))
        elif s["method"] == "sim" and res.get("measures"):
            _table(doc, ["Measure", "Baseline mean", "Solution mean", "p", "Criterion met"],
                   [[x["label"], f"{x['baseline'].get('mean', 0):.3g}", f"{x['solution'].get('mean', 0):.3g}",
                     f"{x.get('p_value', float('nan')):.2g}", str(x.get("criterion_met"))] for x in res["measures"]])
        elif s["method"] == "sej" and res.get("classical"):
            _table(doc, ["Expert", "Calibration", "Information", "Weight"],
                   [[k, f"{v['calibration']:.3f}", f"{v['information']:.3f}", f"{v['weight']:.3f}"] for k, v in res["classical"]["experts"].items()])
        elif s["method"] == "orc" and m.get("occurrences"):
            doc.add_paragraph(f"{len(m['occurrences'])} occurrences classified (ERC / RAT).")
        elif s["method"] == "gsn" and m.get("nodes"):
            doc.add_paragraph(f"Safety argument with {sum(1 for n in m['nodes'] if n['type'] == 'goal')} goals and "
                              f"{sum(1 for n in m['nodes'] if n['type'] == 'solution')} solutions.")
        elif s["method"] == "inv" and m.get("occurrence"):
            doc.add_paragraph(f"Investigation {m['occurrence'].get('ref', '')}: {m['occurrence'].get('summary', '')}")
        elif s["method"] == "wildlife" and res.get("species"):
            doc.add_paragraph(f"Wildlife strike risk by species ({res['years'][0]}–{res['years'][-1]}): {res['total_strikes']:.0f} strikes, "
                              f"{res['rate_per_10k']:.2f} per 10,000 movements.")
            _table(doc, ["Rank", "Species", "Strikes/yr", "Damaging %", "L × S", "Risk"],
                   [[str(x["rank"]), f"{x['common']} ({x['scientific']})", f"{x['per_year']:.1f}",
                     "—" if x["damage_pct"] is None else f"{x['damage_pct']:.1f}", f"{x['likelihood']} × {x['severity']} = {x['score']}",
                     x["risk"]] for x in res["species"]])
        elif m.get("rows"):
            doc.add_paragraph(f"{len(m['rows'])} worksheet rows recorded (see NAVRAP for the full worksheet).")
        else:
            doc.add_paragraph("See NAVRAP for the full model.")
    doc.add_heading("6. Residual risk, actions and conditions", 1)
    _table(doc, ["Ref", "Action", "Owner", "Due", "Status"],
           [[x["ref"], x["text"], x["owner"], x["due_date"] or "", x["status"]] for x in a["actions"]] or [["—"] * 5])
    doc.add_heading("7. Review and approval history", 1)
    _table(doc, ["When", "Action", "From → to", "By", "Comment"],
           [[x["at"][:16], x["action"], f"{x['from_status']} → {x['to_status']}", x["username"], x["comment"]] for x in a["approvals"]] or [["—"] * 5])
    f = doc.add_paragraph("Numbers produced by NAVRAP calculation engines are order-of-magnitude estimates; "
                          "their inputs and sources are recorded in the studies.")
    f.alignment = WD_ALIGN_PARAGRAPH.LEFT
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()
